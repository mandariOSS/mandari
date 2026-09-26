# SPDX-License-Identifier: AGPL-3.0-or-later
"""
Rechte- und Mandantengrenzen im Work-Portal: Matrix aller schreibenden Endpunkte.

Jeder Work-Endpunkt, der POST oder DELETE annimmt, steht mit gültigen Nutzdaten in ``CASES`` (je Formular-/HTMX-Aktion
ein eigener Fall, ID ``<url_name>[:<aktion>]``). Geprüft wird:

- ``test_mitglied_ohne_rechte``: aktives Mitglied von Org A, dessen Rolle keine Berechtigungen hat (kein Gast)
  → 403 oder 404.
- ``test_fremde_organisation``: Administrator von Org B unter dem Slug von Org A → 403 (keine Mitgliedschaft).
- ``test_fremde_objekte_im_pfad``: derselbe Administrator unter dem eigenen Slug, aber mit Objekt-IDs aus Org A im Pfad
  → 403 oder 404. Unterobjekte des Pfad-Objekts in den Nutzdaten (TOP, Protokolleintrag, Checklistenpunkt …) stammen
  ebenfalls aus Org A, alle übrigen Referenzen (Mitglieder, Ordner, Themen …) aus Org B – so scheitert die Anfrage
  nur an der Organisationsgrenze des Pfad-Objekts.
- ``test_fremde_objekte_in_nutzdaten``: Pfad-Objekte aus Org B, einzelne IDs in den Nutzdaten aus Org A – das
  Zielobjekt der Aktion oder eine Referenz, die gespeichert würde (ID ``<fall>:<variante>``). Hier zählt die Grenze
  selbst: Keine Zeile von Org A wird geändert oder gelöscht, und keine neue oder geänderte Zeile verweist auf Org A.
  Die Aktion auf den eigenen Objekten darf dabei ohne den fremden Verweis ausgeführt oder abgelehnt werden (kein 5xx).
  Optionale Vorbelegungen und Mehrfachauswahl-Listen, die fremde IDs verwerfen, sind nicht Teil der Matrix.
- ``test_selbstbedienung_fremdes_objekt``: Selbstbedienungs-Endpunkte mit Objekt-ID verweigern einem Basis-Mitglied
  das Objekt eines anderen Mitglieds derselben Organisation → 403 oder 404.

In allen übrigen Fällen darf sich die Datenbank nicht ändern: Nach dem Aufruf werden alle Zeilen der Apps aus
``WATCHED_APPS`` mit dem Stand nach dem Aufbau verglichen. Ein 400 oder eine Weiterleitung zählt nicht als Abweisung.
Bei Aufrufen unter dem Slug von Org B vermerkt die Fehlermeldung zusätzlich, wenn eine neue oder geänderte Zeile auf
ein Objekt von Org A verweist.

Von der Prüfung ohne Rechte ausgenommen sind die Selbstbedienungs-Endpunkte in ``SELBSTBEDIENUNG``: Sie ändern nur
eigene Daten des angemeldeten Mitglieds. RIS-Objekte (Sitzung, TOP, Vorlage, Datei) gelten als Objekte von Org A, wenn
ihre Körperschaft nur mit Org A verknüpft ist.

Die Testdaten werden modulweit angelegt; jeder Fall läuft in einer Transaktion, die zurückgerollt wird.
"""

from __future__ import annotations

import copy
import io
import json
import uuid
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, time, timedelta
from typing import Any, cast

import docx
import pytest
from django.apps import apps as django_apps
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db.models import Model
from django.http import HttpResponse
from django.test import Client, override_settings
from django.urls import reverse
from django.utils import timezone

import apps.work.urls as work_urls
from apps.accounts.models import User
from apps.common.models import AISettings, SiteSettings
from apps.common.tests.factories import DEFAULT_PASSWORD, MembershipFactory, OrganizationFactory, UserFactory
from apps.tenants.models import (
    AdministrationContact,
    CouncilParty,
    Membership,
    Organization,
    PartyGroup,
    Permission,
    Role,
    Topic,
    UserInvitation,
)
from apps.work.faction.models import (
    FactionAgendaItem,
    FactionAgendaItemAttachment,
    FactionAttendance,
    FactionMeeting,
    FactionMeetingException,
    FactionMeetingSchedule,
    FactionProtocolEntry,
    FactionSuspensionRule,
)
from apps.work.meetings import services as meeting_services
from apps.work.meetings.models import AgendaItemNote, FileAnnotation
from apps.work.motions.models import (
    DocumentFolder,
    FolderGuestShare,
    Motion,
    MotionApproval,
    MotionChecklistItem,
    MotionComment,
    MotionRevision,
    MotionShare,
    MotionTemplate,
    MotionType,
    OrganizationLetterhead,
)
from apps.work.notifications.models import Notification, NotificationType
from apps.work.organization.models import DataExport, MemberAbsence, MemberChangeRequest
from apps.work.support.models import KnowledgeBaseArticle, KnowledgeBaseCategory, SupportTicket
from apps.work.tasks.models import Task, TaskAttachment, TaskChecklistItem, TaskLabel
from insight_core.models import (
    OParlAgendaItem,
    OParlBody,
    OParlConsultation,
    OParlFile,
    OParlMeeting,
    OParlOrganization,
    OParlPaper,
    OParlPerson,
    OParlSource,
)

pytestmark = pytest.mark.django_db

SLUGS = {"A": "matrix-fraktion-a", "B": "matrix-fraktion-b"}

#: Apps, deren Tabellen vor/nach jedem Aufruf verglichen werden
WATCHED_APPS = frozenset({"accounts", "common", "insight_core", "session", "tenants", "work"})

#: Selbstbedienung: ändert ausschließlich eigene Daten des angemeldeten Mitglieds, daher nicht Teil der Prüfung
#: „ohne Rechte“ (nur fremde Organisation, fremde Objekt-IDs und fremde Objekte derselben Organisation).
SELBSTBEDIENUNG: dict[str, str] = {
    "profile": "eigenes Profil (Name, Telefon, Profilbild, Kalender-Feed)",
    "security": "eigene Kontosicherheit (Passwort, 2FA, Sitzungen, vertrauenswürdige Geräte)",
    "profile_notifications": "eigene Benachrichtigungseinstellungen",
    "notification_preferences": "Weiterleitung zu den eigenen Benachrichtigungseinstellungen",
    "notifications_mark_all_read": "eigene Benachrichtigungen als gelesen markieren",
    "notification_mark_read": "eigene Benachrichtigung als gelesen markieren",
    "profile_absence": "eigene Abwesenheit eintragen oder stornieren",
    "profile_requests": "eigenen Änderungsantrag einreichen oder zurückziehen",
    "profile_data": "eigener Datenexport und Löschung der eigenen Mitgliedschaft",
    "export_delete": "eigenen Datenexport löschen",
    "profile_visibility": "Sichtbarkeit des eigenen Profils",
    "profile_committees": "eigene Gremien und Fachgebiete",
    "session_invitation_respond": "eigene Ladung der Verwaltung beantworten",
}

#: Aktionen auf Selbstbedienungs-Seiten, die Daten anderer ändern und deshalb voll geprüft werden
PRUEFAKTIONEN = frozenset({"profile_requests:approve_request", "profile_requests:reject_request"})

#: Schreibende Work-URLs ohne Org-Kontext
NICHT_IN_MATRIX: dict[str, str] = {
    "accept_invitation": "öffentliche Annahme einer Einladung per Token, ohne Org-Slug",
}

#: Unterobjekte von Pfad-Objekten: bei fremden Pfad-Objekten stammen sie aus derselben Organisation wie der Pfad
SUB_OBJECTS = frozenset(
    {
        "checklist",
        "fitem",
        "fitem2",
        "fproposal",
        "fentry",
        "fattendance",
        "fattachment",
        "task_checklist",
        "task_checklist2",
        "task_attachment",
    }
)

HEUTE = date.today()
IN_60_TAGEN = (HEUTE + timedelta(days=60)).isoformat()
IN_65_TAGEN = (HEUTE + timedelta(days=65)).isoformat()
NEUES_PASSWORT = "Neues-Passwort-2026-lang-genug!"

PDF_BYTES = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"
CSV_BYTES = b"titel;status\nImportierte Aufgabe;todo\n"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

Snapshot = dict[str, dict[Any, tuple[Any, ...]]]


# =============================================================================
# Fälle
# =============================================================================


@dataclass(frozen=True)
class Case:
    """Ein schreibender Aufruf: URL-Name, Pfad-Objekte, Methode und gültige Nutzdaten."""

    url_name: str
    action: str = ""
    #: URL-Parameter → Objektschlüssel; ``=wert`` setzt den Wert wörtlich
    path: dict[str, str] = field(default_factory=dict)
    #: Nutzdaten mit Platzhaltern ``{schlüssel}`` (Objekt der Anfrage-Org) und ``{a_schlüssel}`` (Objekt von Org A)
    data: dict[str, Any] = field(default_factory=dict)
    method: str = "post"  # post | json | delete
    #: Formularfeld → Dateiname einer Beispieldatei
    files: dict[str, str] = field(default_factory=dict)
    #: fremde IDs aus Org A in den Nutzdaten: Variante → überschriebene Felder
    foreign: dict[str, dict[str, Any]] = field(default_factory=dict)
    #: Pfad-Objekte aus Org A unter dem Slug von Org B prüfen
    idor: bool = True
    #: Selbstbedienung: Objekt eines anderen Mitglieds derselben Organisation prüfen
    same_org: bool = False

    @property
    def name(self) -> str:
        return f"{self.url_name}:{self.action}" if self.action else self.url_name

    @property
    def self_service(self) -> bool:
        return self.url_name in SELBSTBEDIENUNG and self.name not in PRUEFAKTIONEN

    @property
    def has_path_objects(self) -> bool:
        return any(not key.startswith("=") for key in self.path.values())


MOTION = {"motion_id": "motion"}
FITEM = {"meeting_id": "fmeeting", "item_id": "fitem"}
TASK = {"task_id": "task"}
RIS_ITEM = {"meeting_id": "ris_meeting", "item_id": "ris_item"}

LINK = {"title": "Link", "url": "https://example.org/dokument", "document_type": "link"}
ANMERKUNG = {"content": "Anmerkung", "page": 1}
EINREICHUNG = {
    "title": "Antrag Radweg",
    "application_type": "motion",
    "resolution_proposal": "Der Rat beschließt den Radweg.",
    "justification": "Mehr Sicherheit.",
    "confirm": "on",
}
AUFGABE_FORMULAR = {
    "title": "Aufgabe geändert",
    "status": "in_progress",
    "priority": "high",
    "visibility": "organization",
}
REDNER = {"entry_type": "speech", "speaker": "{a_member}"}
ZUSTAENDIG = {"entry_type": "action", "action_assignee": "{a_member}"}


def _fraktion(
    action: str,
    foreign: dict[str, dict[str, Any]] | None = None,
    meeting: str = "fmeeting",
    **data: Any,
) -> Case:
    return Case(
        "faction_action", action, path={"meeting_id": meeting}, data={"action": action, **data}, foreign=foreign or {}
    )


def _top(
    action: str,
    foreign: dict[str, dict[str, Any]] | None = None,
    files: dict[str, str] | None = None,
    **data: Any,
) -> Case:
    return Case(
        "faction_item_panel_action",
        action,
        path=FITEM,
        data={"action": action, **data},
        files=files or {},
        foreign=foreign or {},
    )


def _aufgabe(
    action: str,
    foreign: dict[str, dict[str, Any]] | None = None,
    files: dict[str, str] | None = None,
    **data: Any,
) -> Case:
    return Case(
        "task_panel_action",
        action,
        path=TASK,
        data={"action": action, **data},
        files=files or {},
        foreign=foreign or {},
    )


def _reihe(action: str, foreign: dict[str, dict[str, Any]] | None = None, **data: Any) -> Case:
    return Case("organization_faction_settings", action, data={"section": action, **data}, foreign=foreign or {})


def _mitglied(action: str, foreign: dict[str, dict[str, Any]] | None = None, **data: Any) -> Case:
    return Case(
        "member_detail", action, path={"member_id": "member"}, data={"action": action, **data}, foreign=foreign or {}
    )


def _rat(action: str, foreign: dict[str, dict[str, Any]] | None = None, **data: Any) -> Case:
    return Case("council_parties", action, data={"action": action, **data}, foreign=foreign or {})


def _sicherheit(action: str, **data: Any) -> Case:
    return Case("security", action, data={"action": action, **data})


CASES: list[Case] = [
    # --- Sitzungsvorbereitung (RIS) ---
    # Ohne Sitzungsdienst-Daten: die ID ist unbekannt, die Grenze liegt in der Zuordnung Person ↔ Ladung
    Case("session_invitation_respond", path={"recipient_id": "unbekannt"}, data={"action": "confirm"}, idor=False),
    Case(
        "meeting_prepare",
        "mark_prepared",
        path={"meeting_id": "ris_meeting"},
        data={"action": "mark_prepared", "notes": "Vorbereitet"},
    ),
    Case("meeting_prepare", "notizen", path={"meeting_id": "ris_meeting"}, data={"notes": "Notizen"}, method="json"),
    Case("meeting_position_api", path=RIS_ITEM, data={"position": "for", "reasoning": "Begründung"}, method="json"),
    Case("meeting_notes_api", path=RIS_ITEM, data={"content": "Beitrag", "visibility": "organization"}, method="json"),
    Case("meeting_note_delete", path={"note_id": "note"}, method="delete"),
    Case("meeting_private_note_api", path=RIS_ITEM, data={"content": "Private Notiz"}, method="json"),
    Case(
        "meeting_speech_api",
        "speichern",
        path=RIS_ITEM,
        data={"title": "Rede", "content": "<p>Rede</p>", "is_shared": True},
        method="json",
    ),
    Case("meeting_speech_api", "loeschen", path=RIS_ITEM, method="delete"),
    Case("meeting_supplementary_api", "link", path=RIS_ITEM, data=LINK, method="json"),
    Case("meeting_supplementary_api", "upload", path=RIS_ITEM, data={"title": "Anlage"}, files={"file": "anlage.pdf"}),
    Case("meeting_supplementary_delete", path={"doc_id": "supp_doc"}, method="delete"),
    Case("meeting_documents_api", path=RIS_ITEM, data=LINK, method="json"),
    Case("meeting_document_delete", path={"link_id": "supp_doc"}, method="delete"),
    Case(
        "meeting_file_annotations",
        "anlage",
        path={"anchor_type": "=doc", "file_id": "supp_doc"},
        data=ANMERKUNG,
        method="json",
    ),
    Case(
        "meeting_file_annotations",
        "ris_datei",
        path={"anchor_type": "=oparl", "file_id": "oparl_file"},
        data=ANMERKUNG,
        method="json",
    ),
    Case("meeting_file_annotation_delete", path={"annotation_id": "annotation"}, method="delete"),
    Case(
        "paper_comments_api",
        path={"paper_id": "paper"},
        data={"content": "Kommentar", "visibility": "organization"},
        method="json",
    ),
    Case("paper_comment_delete", path={"comment_id": "paper_comment"}, method="delete"),
    # --- Dokumente ---
    Case(
        "document_create",
        data={"title": "Neues Dokument", "summary": "Kurzfassung", "folder": "{folder}"},
        foreign={"ordner": {"folder": "{a_folder}"}},
    ),
    Case("document_import", data={"visibility": "private"}, files={"import_files": "antrag.docx"}),
    Case(
        "document_folder_create",
        data={"name": "Neuer Ordner", "parent": "{folder}"},
        foreign={"ordner": {"parent": "{a_folder}"}},
    ),
    Case(
        "document_folder_update",
        path={"folder_id": "folder"},
        data={"name": "Ablage neu", "color": "green", "parent": "{folder2}"},
        foreign={"ordner": {"parent": "{a_folder}"}},
    ),
    Case("document_folder_delete", path={"folder_id": "folder2"}),
    Case("document_folder_share", path={"folder_id": "folder"}, data={"email": "{guest_email}", "level": "comment"}),
    Case("document_folder_share_remove", path={"share_id": "folder_share"}),
    Case(
        "document_move_to_folder",
        data={"folder": "{folder2}", "motion_ids": ["{motion}"]},
        foreign={"ordner": {"folder": "{a_folder}"}, "dokument": {"motion_ids": ["{a_motion}"]}},
    ),
    Case("document_ai", data={"action": "title", "text": "Radweg an der Hauptstraße"}),
    Case(
        "document_editor",
        "save",
        path=MOTION,
        data={"action": "save", "title": "Neuer Titel", "content": "<p>Neuer Inhalt</p>"},
        foreign={"dokumenttyp": {"document_type_id": "{a_doctype}"}, "briefkopf": {"letterhead_id": "{a_letterhead}"}},
    ),
    Case("document_editor", "delete", path=MOTION, data={"action": "delete"}),
    Case(
        "document_editor",
        "formular",
        path=MOTION,
        data={"action": "form", "motion_type": "motion", "title": "Titel per Formular", "tags": "[]"},
    ),
    Case(
        "document_share_update",
        path=MOTION,
        data={"visibility": "shared", "add_user_email": "{member_email}", "level": "comment"},
    ),
    Case("document_share_remove", path={"share_id": "share"}),
    Case("document_status", path=MOTION, data={"status": "internal_review"}),
    Case(
        "document_meta",
        "set_responsible",
        path=MOTION,
        data={"action": "set_responsible", "responsible": "{member}"},
        foreign={"mitglied": {"responsible": "{a_member}"}},
    ),
    Case(
        "document_meta",
        "set_contributors",
        path=MOTION,
        data={"action": "set_contributors", "contributors": ["{member}"]},
    ),
    Case("document_meta", "set_topics", path=MOTION, data={"action": "set_topics", "topics": ["{topic}"]}),
    Case(
        "document_meta",
        "set_folder",
        path=MOTION,
        data={"action": "set_folder", "folder": "{folder2}"},
        foreign={"ordner": {"folder": "{a_folder}"}},
    ),
    Case("document_meta", "set_due_date", path=MOTION, data={"action": "set_due_date", "due_date": IN_60_TAGEN}),
    Case("document_checklist", "add", path=MOTION, data={"action": "add", "title": "Neuer Punkt"}),
    Case(
        "document_checklist",
        "toggle",
        path=MOTION,
        data={"action": "toggle", "item_id": "{checklist}"},
        foreign={"punkt": {"item_id": "{a_checklist}"}},
    ),
    Case(
        "document_checklist",
        "delete",
        path=MOTION,
        data={"action": "delete", "item_id": "{checklist}"},
        foreign={"punkt": {"item_id": "{a_checklist}"}},
    ),
    Case(
        "document_approval_request",
        path=MOTION,
        data={"approver": "{member}", "approval_type": "council"},
        foreign={"mitglied": {"approver": "{a_member}"}},
    ),
    Case(
        "document_approval_decide",
        path={"motion_id": "motion", "approval_id": "approval"},
        data={"decision": "approve", "comment": "Einverstanden"},
    ),
    Case(
        "document_comment",
        path=MOTION,
        data={"content": "Neuer Kommentar"},
        foreign={"antwort": {"parent": "{a_comment}"}},
    ),
    Case("document_comment_resolve", path={"motion_id": "motion", "comment_id": "comment"}),
    Case("document_upload", path=MOTION, files={"file": "anlage.pdf"}),
    Case("document_submit_ris", path=MOTION, data=EINREICHUNG),
    Case("document_empty_trash"),
    Case("document_restore", path={"motion_id": "motion_trash"}),
    Case("document_permanent_delete", path={"motion_id": "motion_trash"}),
    Case("document_revision_restore", path={"motion_id": "motion", "revision_id": "revision"}),
    # --- Dokument-Einstellungen ---
    Case("document_type_create", data={"name": "Anfrage", "slug": "anfrage-neu"}),
    Case("document_type_edit", path={"type_id": "doctype"}, data={"name": "Antrag geändert", "slug": "antrag"}),
    Case("document_type_delete", path={"type_id": "doctype"}),
    Case("document_topic_list", data={"name": "Neues Thema", "color": "green"}),
    Case("document_topic_update", path={"topic_id": "topic"}, data={"name": "Thema neu", "color": "red"}),
    Case("document_topic_delete", path={"topic_id": "topic"}),
    Case(
        "document_template_create",
        data={"name": "Neue Vorlage", "motion_type": "{doctype}", "letterhead": "{letterhead}", "is_active": "on"},
    ),
    Case(
        "document_template_edit",
        path={"template_id": "template"},
        data={"name": "Vorlage geändert", "motion_type": "{doctype}", "is_active": "on"},
    ),
    Case("document_template_delete", path={"template_id": "template"}),
    Case("document_letterhead_create", data={"name": "Neuer Briefkopf", "kind": "generated"}),
    Case(
        "document_letterhead_edit",
        path={"letterhead_id": "letterhead"},
        data={"name": "Briefkopf geändert", "kind": "generated"},
    ),
    Case("document_letterhead_delete", path={"letterhead_id": "letterhead_free"}),
    # --- Fraktionssitzungen ---
    Case("faction", data={"title": "Neue Sitzung", "start_date": IN_60_TAGEN, "start_time": "18:00"}),
    Case("faction_settings"),
    _fraktion("start"),
    _fraktion("end", meeting="fmeeting_ongoing"),
    _fraktion("cancel"),
    _fraktion("delete"),
    _fraktion("update_status", status="invited"),
    _fraktion("update", title="Sitzung geändert", start_date=IN_65_TAGEN, start_time="19:00"),
    _fraktion("invite"),
    _fraktion("release_invitations"),
    _fraktion("add_item", title="Neuer TOP", visibility="public", foreign={"eltern_top": {"parent_id": "{a_fitem}"}}),
    _fraktion("edit_item", item_id="{fitem}", title="TOP geändert", foreign={"top": {"item_id": "{a_fitem}"}}),
    _fraktion("delete_item", item_id="{fitem}", foreign={"top": {"item_id": "{a_fitem}"}}),
    _fraktion("move_item", item_id="{fitem2}", direction="up", foreign={"top": {"item_id": "{a_fitem2}"}}),
    _fraktion(
        "add_entry",
        entry_type="note",
        content="Eintrag",
        agenda_item_id="{fitem}",
        foreign={"redner": REDNER, "zustaendig": ZUSTAENDIG},
    ),
    _fraktion(
        "edit_entry",
        entry_id="{fentry}",
        content="Eintrag geändert",
        foreign={"eintrag": {"entry_id": "{a_fentry}"}, "redner": {"speaker": "{a_member}"}},
    ),
    _fraktion("delete_entry", entry_id="{fentry}", foreign={"eintrag": {"entry_id": "{a_fentry}"}}),
    _fraktion(
        "record_decision",
        agenda_item_id="{fitem}",
        votes_yes="5",
        votes_no="1",
        votes_abstain="0",
        result="accepted",
        foreign={"top": {"agenda_item_id": "{a_fitem}"}},
    ),
    _fraktion("approve_protocol", meeting="fmeeting_completed"),
    _fraktion("respond", status="confirmed"),
    _fraktion("check_in", attendance_id="{fattendance}", foreign={"teilnahme": {"attendance_id": "{a_fattendance}"}}),
    _fraktion("check_out", attendance_id="{fattendance}", foreign={"teilnahme": {"attendance_id": "{a_fattendance}"}}),
    _fraktion(
        "add_attendee",
        attendee_type="guest",
        guest_name="Gast",
        status="present",
        foreign={"mitglied": {"attendee_type": "member", "membership_id": "{a_member}"}},
    ),
    _fraktion(
        "set_participation",
        attendance_id="{fattendance}",
        participation_type="online",
        foreign={"teilnahme": {"attendance_id": "{a_fattendance}"}},
    ),
    _fraktion("confirm_attendance", meeting="fmeeting_completed"),
    _fraktion("propose", title="Vorschlag", description="Bitte aufnehmen", visibility="public"),
    _fraktion("accept_proposal", item_id="{fproposal}", foreign={"vorschlag": {"item_id": "{a_fproposal}"}}),
    _fraktion(
        "reject_proposal",
        item_id="{fproposal}",
        reason="Nicht jetzt",
        foreign={"vorschlag": {"item_id": "{a_fproposal}"}},
    ),
    _fraktion("create_task", entry_id="{fentry}", foreign={"eintrag": {"entry_id": "{a_fentry}"}}),
    _top("update", title="TOP neu", description="Beschreibung", visibility="public"),
    _top("add_entry", entry_type="note", content="Eintrag", foreign={"redner": REDNER, "zustaendig": ZUSTAENDIG}),
    _top(
        "edit_entry",
        entry_id="{fentry}",
        content="Eintrag geändert",
        foreign={"eintrag": {"entry_id": "{a_fentry}"}, "redner": {"speaker": "{a_member}"}},
    ),
    _top("delete_entry", entry_id="{fentry}", foreign={"eintrag": {"entry_id": "{a_fentry}"}}),
    _top("record_decision", votes_yes="4", votes_no="0", votes_abstain="1", result="accepted"),
    _top("clear_decision"),
    _top("upload_attachment", files={"file": "anlage.pdf"}),
    _top(
        "delete_attachment",
        attachment_id="{fattachment}",
        foreign={"anhang": {"attachment_id": "{a_fattachment}"}},
    ),
    _top("link_motion", motion_id="{motion2}", foreign={"dokument": {"motion_id": "{a_motion2}"}}),
    _top("unlink_motion", motion_id="{motion}"),
    _top("link_paper", paper_id="{paper2}"),
    _top("unlink_paper", paper_id="{paper}"),
    _top(
        "create_task",
        title="Aufgabe aus TOP",
        assigned_to="{member}",
        due_date=IN_60_TAGEN,
        foreign={"zustaendig": {"assigned_to": "{a_member}"}},
    ),
    _top("add_link", link_label="Quelle", link_url="https://example.org/quelle"),
    _top("remove_link", link_index="0"),
    # --- Aufgaben ---
    Case(
        "tasks_api",
        "move",
        data={"action": "move", "task_id": "{task}", "status": "in_progress", "position": 0},
        method="json",
        foreign={"aufgabe": {"task_id": "{a_task}"}},
    ),
    Case(
        "tasks_api",
        "quick_add",
        data={"action": "quick_add", "title": "Schnelle Aufgabe", "status": "todo", "priority": "medium"},
    ),
    Case(
        "tasks_api",
        "update_status",
        data={"action": "update_status", "task_id": "{task}", "status": "done"},
        foreign={"aufgabe": {"task_id": "{a_task}"}},
    ),
    Case(
        "tasks_api",
        "toggle_complete",
        data={"action": "toggle_complete", "task_id": "{task}"},
        foreign={"aufgabe": {"task_id": "{a_task}"}},
    ),
    Case("tasks_import", data={"entry_ids[]": ["{fentry}"]}),
    Case("tasks_import_file", files={"file": "aufgaben.csv"}),
    Case(
        "task_create",
        data={
            "title": "Neue Aufgabe",
            "visibility": "organization",
            "priority": "medium",
            "status": "todo",
            "tags": "[]",
            "related_motion": "{motion}",
        },
    ),
    Case("task_labels", data={"name": "Neu", "color": "green"}),
    Case("task_label_delete", path={"label_id": "label"}, method="delete"),
    Case(
        "task_share",
        path=TASK,
        data={"visibility": "shared", "share_with[]": ["{member}"]},
        foreign={"mitglied": {"share_with[]": ["{a_member}"]}},
    ),
    Case("task_panel_action", "update", path=TASK, data={"action": "update", **AUFGABE_FORMULAR}),
    Case("task_panel_action", "save", path=TASK, data={"action": "save", **AUFGABE_FORMULAR}),
    _aufgabe("add_comment", content="Kommentar"),
    _aufgabe("toggle_complete"),
    _aufgabe("delete"),
    _aufgabe("upload_attachment", files={"file": "anlage.pdf"}),
    _aufgabe(
        "delete_attachment",
        attachment_id="{task_attachment}",
        foreign={"anhang": {"attachment_id": "{a_task_attachment}"}},
    ),
    _aufgabe("add_checklist_item", title="Schritt"),
    _aufgabe(
        "toggle_checklist_item",
        item_id="{task_checklist}",
        foreign={"punkt": {"item_id": "{a_task_checklist}"}},
    ),
    _aufgabe(
        "delete_checklist_item",
        item_id="{task_checklist}",
        foreign={"punkt": {"item_id": "{a_task_checklist}"}},
    ),
    _aufgabe("toggle_label", label_id="{label}", foreign={"label": {"label_id": "{a_label}"}}),
    _aufgabe(
        "reorder_checklist",
        order='["{task_checklist2}", "{task_checklist}"]',
        foreign={"punkt": {"order": '["{a_task_checklist}"]'}},
    ),
    # --- Organisation ---
    Case(
        "organization",
        "update_general",
        data={
            "action": "update_general",
            "name": "Fraktion umbenannt",
            "description": "Neu",
            "primary_color": "#123456",
        },
    ),
    Case(
        "organization",
        "update_contact",
        data={
            "action": "update_contact",
            "contact_email": "kontakt@example.org",
            "contact_phone": "0123 456",
            "website": "https://example.org",
            "address": "Rathausplatz 1",
        },
    ),
    Case("organization", "update_parties", data={"action": "update_parties", "parties": ["{party_group}"]}),
    Case("organization", "update_security", data={"action": "update_security", "require_2fa": "on"}),
    Case(
        "organization_faction_settings",
        "speichern",
        data={"auto_create_approval_item": "on", "invitation_lead_hours": "48"},
    ),
    _reihe("add_schedule", name="Reihe", time="18:00", weekday="2", duration_minutes="90", recurrence="weekly"),
    _reihe("toggle_schedule", schedule_id="{schedule}", foreign={"reihe": {"schedule_id": "{a_schedule}"}}),
    _reihe("delete_schedule", schedule_id="{schedule}", foreign={"reihe": {"schedule_id": "{a_schedule}"}}),
    _reihe(
        "add_exception",
        schedule_id="{schedule}",
        original_date=IN_60_TAGEN,
        end_date=IN_65_TAGEN,
        reason="Ferien",
        foreign={"reihe": {"schedule_id": "{a_schedule}"}},
    ),
    _reihe("delete_exception", exception_id="{exception}", foreign={"ausnahme": {"exception_id": "{a_exception}"}}),
    _reihe(
        "add_rule",
        schedule_id="{schedule}",
        ris_organization_id="{committee2}",
        foreign={"reihe": {"schedule_id": "{a_schedule}"}},
    ),
    _reihe("delete_rule", rule_id="{rule}", foreign={"regel": {"rule_id": "{a_rule}"}}),
    Case(
        "organization_api_settings",
        "api_save",
        data={
            "section": "api_save",
            "api_enabled": "on",
            "api_past_days": "30",
            "api_future_days": "90",
            "api_cache_seconds": "300",
        },
    ),
    Case("organization_api_settings", "api_regenerate", data={"section": "api_regenerate"}),
    Case("organization_ris_settings", "connect", data={"token": "ungueltiges-token"}),
    Case("organization_ris_settings", "disconnect", data={"action": "disconnect"}),
    Case("organization_email_settings", "save", data={"action": "save", "mail_sender_mode": "mandari"}),
    Case("organization_email_settings", "send_test", data={"action": "send_test"}),
    Case("member_invite", data={"email": "neu@matrix.example.org", "roles": ["{role}"], "message": "Willkommen"}),
    Case(
        "guest_invite",
        data={
            "email": "gast-neu@matrix.example.org",
            "share_level": "view",
            "documents": ["{motion}"],
            "folders": ["{folder}"],
        },
    ),
    _mitglied("update_committees", committees=["{committee}"]),
    _mitglied("update_expertise", expertise_topics=["{topic}"]),
    _mitglied("update_roles", roles=["{role}"]),
    _mitglied("update_permissions", individual_permissions=["dashboard.view"]),
    _mitglied("deactivate"),
    _mitglied("reactivate"),
    _mitglied("remove"),
    _mitglied("transfer_ownership"),
    _mitglied("link_oparl_person", oparl_person_id="{person}", foreign={"person": {"oparl_person_id": "{a_person}"}}),
    _mitglied("unlink_oparl_person"),
    _mitglied("apply_suggestions"),
    _mitglied("update_sworn_in", is_sworn_in="1"),
    Case("guest_access_resend", path={"member_id": "guest"}, data={"message": "Ihr Zugang"}),
    Case("invitation_resend", path={"invitation_id": "invitation"}),
    Case("invitation_cancel", path={"invitation_id": "invitation"}),
    Case(
        "role_create",
        data={"name": "Neue Rolle", "color": "#123456", "priority": "40", "permissions": ["motions.view"]},
    ),
    Case("role_edit", path={"role_id": "role"}, data={"name": "Sonderrolle geändert", "permissions": ["motions.view"]}),
    Case("role_delete", path={"role_id": "role"}),
    Case("role_reset", path={"role_id": "role_std"}),
    Case("roles_restore_defaults"),
    Case(
        "organization_registration",
        data={
            "registration_enabled": "on",
            "registration_email_domains": "example.org",
            "registration_default_role": "{role}",
        },
    ),
    Case("member_approve", path={"membership_id": "pending"}),
    Case("member_reject", path={"membership_id": "pending"}, data={"reason": "Keine Zugehörigkeit"}),
    _rat("update_coalition", coalition_name="Koalition"),
    _rat("add_admin_contact", contact_label="Ratsbüro", contact_email="rat@example.org"),
    _rat("delete_admin_contact", contact_id="{contact}", foreign={"kontakt": {"contact_id": "{a_contact}"}}),
    _rat("add_party", name="Neue Partei", short_name="NP", color="#123456", is_active="on"),
    _rat(
        "update_party",
        party_id="{party}",
        name="Partei geändert",
        short_name="PG",
        is_active="on",
        foreign={"partei": {"party_id": "{a_party}"}},
    ),
    _rat("delete_party", party_id="{party}", foreign={"partei": {"party_id": "{a_party}"}}),
    # --- Support ---
    Case(
        "support_create",
        data={"subject": "Frage", "description": "Beschreibung", "category": "question", "priority": "normal"},
    ),
    Case("support_detail", "reply", path={"ticket_id": "ticket"}, data={"action": "reply", "content": "Antwort"}),
    Case("support_detail", "close", path={"ticket_id": "ticket"}, data={"action": "close"}),
    Case("support_detail", "reopen", path={"ticket_id": "ticket_closed"}, data={"action": "reopen"}),
    # Wissensdatenbank ist global (keine Org-Objekte), daher keine Prüfung fremder IDs
    Case(
        "kb_article_feedback",
        path={"article_id": "kb_article"},
        data={"is_helpful": "true", "comment": "Hilfreich"},
        idor=False,
    ),
    # --- Selbstbedienung ---
    Case("notification_preferences"),
    Case("notifications_mark_all_read"),
    Case("notification_mark_read", path={"notification_id": "notification"}, same_org=True),
    Case("profile", "update_profile", data={"action": "update_profile", "first_name": "Erika", "last_name": "Muster"}),
    Case("profile", "remove_avatar", data={"action": "remove_avatar"}),
    Case("profile", "regenerate_calendar_feed", data={"action": "regenerate_calendar_feed"}),
    _sicherheit(
        "change_password",
        old_password=DEFAULT_PASSWORD,
        new_password=NEUES_PASSWORT,
        confirm_password=NEUES_PASSWORT,
    ),
    _sicherheit("setup_2fa"),
    _sicherheit("confirm_2fa", code="123456"),
    _sicherheit("disable_2fa", password=DEFAULT_PASSWORD),
    _sicherheit("regenerate_backup_codes", password=DEFAULT_PASSWORD),
    _sicherheit("revoke_session", session_key="unbekannt"),
    _sicherheit("revoke_all_sessions"),
    _sicherheit("remove_trusted_device", device_id="{unbekannt}"),
    Case("profile_notifications", data={"email_enabled": "on", "email_digest": "daily"}),
    Case(
        "profile_absence",
        "create_absence",
        data={"action": "create_absence", "start_date": IN_60_TAGEN, "end_date": IN_65_TAGEN, "reason": "Urlaub"},
    ),
    Case(
        "profile_absence",
        "cancel_absence",
        data={"action": "cancel_absence", "absence_id": "{absence}"},
        same_org=True,
        foreign={"abwesenheit": {"absence_id": "{a_absence}"}},
    ),
    Case(
        "profile_requests",
        "submit_request",
        data={
            "action": "submit_request",
            "request_type": "role_change",
            "reason": "Bitte",
            "requested_roles": ["{role}"],
        },
    ),
    Case(
        "profile_requests",
        "withdraw_request",
        data={"action": "withdraw_request", "request_id": "{change_request}"},
        same_org=True,
        foreign={"antrag": {"request_id": "{a_change_request}"}},
    ),
    Case(
        "profile_requests",
        "approve_request",
        data={"action": "approve_request", "request_id": "{change_request}"},
        foreign={"antrag": {"request_id": "{a_change_request}"}},
    ),
    Case(
        "profile_requests",
        "reject_request",
        data={"action": "reject_request", "request_id": "{change_request}", "decision_comment": "Nein"},
        foreign={"antrag": {"request_id": "{a_change_request}"}},
    ),
    Case("profile_data", "export_data", data={"action": "export_data", "format": "json"}),
    Case("profile_data", "request_deletion", data={"action": "request_deletion", "password": DEFAULT_PASSWORD}),
    Case("export_delete", path={"export_id": "export"}, same_org=True),
    Case("profile_visibility", data={"bio": "Hallo", "show_email": "on", "preferred_contact": "email"}),
    Case("profile_committees", "gremien", data={"committees": ["{committee}"]}),
    Case("profile_committees", "save_expertise", data={"action": "save_expertise", "expertise_topics": ["{topic}"]}),
]


# =============================================================================
# Testdaten
# =============================================================================


@dataclass
class World:
    """Zwei Organisationen mit denselben Objektarten, dazu eingeloggte Clients und der Ausgangsstand der Datenbank."""

    ids: dict[str, dict[str, str]]
    clients: dict[str, Client]
    baseline: Snapshot = field(default_factory=dict)
    #: IDs aller Objekte von Org A – kennzeichnet in Fehlermeldungen gespeicherte Verweise über die Org-Grenze
    org_a_refs: frozenset[str] = frozenset()


def _create_user(email: str) -> User:
    return cast(User, UserFactory(email=email))  # type: ignore[no-untyped-call]


def _create_org(name: str, slug: str, body: OParlBody) -> Organization:
    return cast(Organization, OrganizationFactory(name=name, slug=slug, body=body))  # type: ignore[no-untyped-call]


def _membership(user: User, organization: Organization, roles: list[Role] | None = None, **kwargs: Any) -> Membership:
    membership = MembershipFactory(user=user, organization=organization, roles=roles, **kwargs)  # type: ignore[no-untyped-call]
    return cast(Membership, membership)


def _login(user: User) -> Client:
    client = Client(raise_request_exception=False)
    client.force_login(user)
    return client


def _docx_bytes() -> bytes:
    dokument = docx.Document()
    dokument.add_heading("Antrag: Mehr Bäume", 1)
    dokument.add_paragraph("Der Rat beschließt, Bäume zu pflanzen.")
    puffer = io.BytesIO()
    dokument.save(puffer)
    return puffer.getvalue()


def _upload(name: str) -> SimpleUploadedFile:
    if name.endswith(".csv"):
        return SimpleUploadedFile(name, CSV_BYTES, content_type="text/csv")
    if name.endswith(".docx"):
        return SimpleUploadedFile(name, _docx_bytes(), content_type=DOCX_MIME)
    return SimpleUploadedFile(name, PDF_BYTES, content_type="application/pdf")


class _Builder:
    """Legt die Objekte einer Organisation an; beide Organisationen erhalten dieselben Objektarten."""

    def __init__(self, source: OParlSource, created_permissions: list[str], users: list[User]) -> None:
        self.source = source
        self.created_permissions = created_permissions
        self.users = users

    def user(self, email: str) -> User:
        user = _create_user(email)
        self.users.append(user)
        return user

    def role(self, organization: Organization, name: str, codes: list[str]) -> Role:
        role = Role.objects.create(organization=organization, name=name)
        for code in codes:
            permission, created = Permission.objects.get_or_create(
                codename=code, defaults={"name": code, "category": code.split(".")[0]}
            )
            if created:
                self.created_permissions.append(code)
            role.permissions.add(permission)
        return role

    def organization(self, key: str) -> tuple[Organization, dict[str, str]]:
        tag = key.lower()
        now = timezone.now()
        base = f"https://ris.example.org/matrix-{tag}"
        body = OParlBody.objects.create(external_id=f"{base}/body", source=self.source, name=f"Stadt {key}")
        org = _create_org(f"Matrix-Fraktion {key}", SLUGS[key], body)

        # Mitglieder: Administrator (Eigentümer), Kollegin, Gast, offene Registrierung
        admin_role = Role.objects.filter(organization=org, is_admin=True).first()
        assert admin_role is not None, f"Standardrollen für {org.slug} fehlen"
        admin_user = self.user(f"admin-{tag}@matrix.example.org")
        admin = _membership(admin_user, org, [admin_role])
        org.owner = admin_user
        cast(Any, org).save(update_fields=["owner"])
        basis_role = self.role(org, "Matrix Basis", ["dashboard.view"])
        kollege_user = self.user(f"kollegin-{tag}@matrix.example.org")
        kollege = _membership(kollege_user, org, [basis_role])
        gast_user = self.user(f"gast-{tag}@matrix.example.org")
        gast = _membership(gast_user, org, is_guest=True)
        anfrage = _membership(
            self.user(f"anfrage-{tag}@matrix.example.org"), org, is_active=False, registration_requested_at=now
        )

        # RIS der eigenen Körperschaft
        committee = OParlOrganization.objects.create(
            external_id=f"{base}/organization/1", body=body, name="Bauausschuss"
        )
        committee2 = OParlOrganization.objects.create(external_id=f"{base}/organization/2", body=body, name="Rat")
        person = OParlPerson.objects.create(external_id=f"{base}/person/1", body=body, name="Ratsmitglied")
        ris_meeting = OParlMeeting.objects.create(
            external_id=f"{base}/meeting/1", body=body, name="Rat", start=now + timedelta(days=3)
        )
        ris_item = OParlAgendaItem.objects.create(
            external_id=f"{base}/agenda/1", meeting=ris_meeting, number="1", name="Haushalt", order=1
        )
        ris_item_paper = OParlAgendaItem.objects.create(
            external_id=f"{base}/agenda/2", meeting=ris_meeting, number="2", name="Spielplatz", order=2
        )
        paper = OParlPaper.objects.create(
            external_id=f"{base}/paper/1", body=body, name="Antrag Spielplatz", reference=f"V/{key}/1"
        )
        paper2 = OParlPaper.objects.create(
            external_id=f"{base}/paper/2", body=body, name="Anfrage Radwege", reference=f"V/{key}/2"
        )
        OParlConsultation.objects.create(
            external_id=f"{base}/consultation/1",
            body=body,
            paper=paper,
            agenda_item_external_id=ris_item_paper.external_id,
            meeting_external_id=ris_meeting.external_id,
            role="Entscheidung",
            authoritative=True,
        )
        oparl_file = OParlFile.objects.create(external_id=f"{base}/file/1", body=body, paper=paper, name="Vorlage")

        # Sitzungsvorbereitung: Thread-Notiz, Vorlagen-Kommentar, ergänzendes Dokument, Anmerkung
        meeting_services.create_thread_note(org, ris_meeting, ris_item, admin, {"content": "Notiz"})
        note = AgendaItemNote.objects.get(organization=org, agenda_item=ris_item)
        paper_comment = meeting_services.create_paper_comment(org, paper, admin, {"content": "Kommentar"})
        supp_doc = meeting_services.add_document_link(
            org, ris_meeting, ris_item, admin, {"title": "Link", "url": "https://example.org/link"}
        )
        meeting_services.add_file_annotation(org, admin, {"supplementary_document": supp_doc}, {"content": "Seite 1"})
        annotation = FileAnnotation.objects.get(organization=org)

        # Dokumente und Einstellungen
        folder = DocumentFolder.objects.create(organization=org, name="Ablage", created_by=admin)
        folder2 = DocumentFolder.objects.create(organization=org, name="Archiv", created_by=admin)
        doctype = MotionType.objects.create(organization=org, name="Antrag", slug="antrag")
        letterhead = OrganizationLetterhead.objects.create(organization=org, name="Briefkopf", kind="generated")
        letterhead_free = OrganizationLetterhead.objects.create(organization=org, name="Unbenutzt", kind="generated")
        template = MotionTemplate.objects.create(
            organization=org, name="Vorlage", motion_type=doctype, letterhead=letterhead
        )
        topic = Topic.objects.create(organization=org, name="Verkehr")
        motion = Motion(
            organization=org,
            author=admin,
            responsible=admin,
            title="Antrag Radweg",
            visibility="organization",
            status="draft",
            folder=folder,
        )
        cast(Any, motion).set_content_encrypted("<p>Inhalt</p>")
        motion.save()
        motion2 = Motion.objects.create(
            organization=org, author=admin, title="Anfrage Beleuchtung", visibility="organization", status="draft"
        )
        motion_trash = Motion.objects.create(
            organization=org,
            author=admin,
            title="Alter Entwurf",
            visibility="organization",
            status="deleted",
            deleted_at=now,
        )
        comment = MotionComment.objects.create(motion=motion, author=admin, content="Bitte prüfen")
        share = MotionShare.objects.create(
            motion=motion, scope="user", user=kollege_user, level="view", created_by=admin_user
        )
        folder_share = FolderGuestShare.objects.create(
            folder=folder, user=gast_user, level="view", created_by=admin_user
        )
        approval = MotionApproval.objects.create(motion=motion, approver=admin, approval_type="chair")
        checklist = MotionChecklistItem.objects.create(motion=motion, title="Prüfen", position=0)
        revision = MotionRevision(motion=motion, version=1, changed_by=admin, change_summary="Erste Fassung")
        cast(Any, revision).set_content_encrypted("<p>Erste Fassung</p>")
        revision.save()

        # Fraktionssitzung mit Tagesordnung, Protokoll, Teilnahmen und Sitzungsreihe
        fmeeting = FactionMeeting.objects.create(
            organization=org,
            title="Fraktionssitzung",
            start=now + timedelta(days=2),
            status="planned",
            created_by=admin,
        )
        fitem = FactionAgendaItem.objects.create(
            meeting=fmeeting,
            number="1",
            title="Haushalt",
            visibility="public",
            order=1,
            reference_links=[{"label": "Quelle", "url": "https://example.org"}],
        )
        fitem.related_motions.add(motion)
        fitem.related_papers.add(paper)
        fitem2 = FactionAgendaItem.objects.create(
            meeting=fmeeting, number="2", title="Verkehr", visibility="public", order=2
        )
        fproposal = FactionAgendaItem.objects.create(
            meeting=fmeeting,
            title="Vorschlag",
            visibility="public",
            order=3,
            proposal_status="proposed",
            proposed_by=kollege,
        )
        fentry = FactionProtocolEntry(
            meeting=fmeeting,
            agenda_item=fitem,
            entry_type="action",
            created_by=admin,
            action_assignee=kollege,
            order=1,
        )
        cast(Any, fentry).set_content_encrypted("Radweg prüfen")
        cast(Any, fentry).save()
        FactionAttendance.objects.create(meeting=fmeeting, membership=admin, status="invited")
        fattendance = FactionAttendance.objects.create(meeting=fmeeting, membership=kollege, status="invited")
        fattachment = FactionAgendaItemAttachment.objects.create(
            agenda_item=fitem,
            file=_upload("top.pdf"),
            filename="top.pdf",
            mime_type="application/pdf",
            file_size=len(PDF_BYTES),
            uploaded_by=admin,
        )
        fmeeting_ongoing = FactionMeeting.objects.create(
            organization=org, title="Laufende Sitzung", start=now, status="ongoing", created_by=admin
        )
        fmeeting_completed = FactionMeeting.objects.create(
            organization=org,
            title="Beendete Sitzung",
            start=now - timedelta(days=7),
            status="completed",
            created_by=admin,
        )
        FactionAttendance.objects.create(meeting=fmeeting_completed, membership=kollege, status="present")
        schedule = FactionMeetingSchedule.objects.create(
            organization=org, name="Wöchentlich", weekday=1, time=time(18, 0)
        )
        exception = FactionMeetingException.objects.create(
            schedule=schedule, original_date=HEUTE + timedelta(days=30), exception_type="cancelled"
        )
        rule = FactionSuspensionRule.objects.create(schedule=schedule, ris_organization=committee)

        # Aufgaben
        task = Task.objects.create(
            organization=org, title="Aufgabe", created_by=admin, assigned_to=admin, visibility="organization"
        )
        label = TaskLabel.objects.create(organization=org, name="Wichtig", color="red")
        task_checklist = TaskChecklistItem.objects.create(task=task, title="Schritt 1", position=0)
        task_checklist2 = TaskChecklistItem.objects.create(task=task, title="Schritt 2", position=1)
        task_attachment = TaskAttachment.objects.create(
            task=task,
            file=_upload("aufgabe.pdf"),
            filename="aufgabe.pdf",
            mime_type="application/pdf",
            file_size=len(PDF_BYTES),
            uploaded_by=admin,
        )

        # Support, Benachrichtigung und Selbstbedienungs-Objekte der Kollegin
        ticket = SupportTicket(organization=org, subject="Frage", created_by=admin)
        cast(Any, ticket).set_description_encrypted("Wie geht das?")
        ticket.save()
        ticket_closed = SupportTicket(organization=org, subject="Erledigt", created_by=admin, status="closed")
        cast(Any, ticket_closed).set_description_encrypted("Schon gelöst")
        ticket_closed.save()
        notification = Notification.objects.create(
            recipient=kollege, notification_type=NotificationType.TASK_ASSIGNED, title="Hinweis", message="Text"
        )
        absence = MemberAbsence.objects.create(
            organization=org,
            membership=kollege,
            start_date=HEUTE + timedelta(days=10),
            end_date=HEUTE + timedelta(days=12),
        )
        change_request = MemberChangeRequest.objects.create(
            organization=org,
            requester=kollege,
            request_type="role_change",
            reason="Mehr Rechte",
            request_data={"requested_roles": []},
        )
        export = DataExport.objects.create(organization=org, membership=kollege)

        # Einladung, Rollen, Ratsfraktion, Verwaltungskontakt
        invitation = cast(Any, UserInvitation).create_for_organization(
            organization=org, email=f"eingeladen-{tag}@matrix.example.org", invited_by=admin_user
        )
        role = Role.objects.create(organization=org, name="Matrix Sonderrolle")
        role_std = Role.objects.filter(organization=org, name="Parteimitglied").first()
        assert role_std is not None, "Standardrolle 'Parteimitglied' fehlt"
        role_std.description = "angepasst"
        role_std.save(update_fields=["description"])
        party = CouncilParty.objects.create(organization=org, name=f"Partei {key}", short_name=f"P{key}")
        contact = AdministrationContact.objects.create(organization=org, label="Ratsbüro", email="rat@example.org")

        objects: dict[str, Any] = {
            "admin": admin,
            "member": kollege,
            "guest": gast,
            "pending": anfrage,
            "committee": committee,
            "committee2": committee2,
            "person": person,
            "ris_meeting": ris_meeting,
            "ris_item": ris_item,
            "paper": paper,
            "paper2": paper2,
            "oparl_file": oparl_file,
            "note": note,
            "paper_comment": paper_comment,
            "supp_doc": supp_doc,
            "annotation": annotation,
            "folder": folder,
            "folder2": folder2,
            "doctype": doctype,
            "letterhead": letterhead,
            "letterhead_free": letterhead_free,
            "template": template,
            "topic": topic,
            "motion": motion,
            "motion2": motion2,
            "motion_trash": motion_trash,
            "comment": comment,
            "share": share,
            "folder_share": folder_share,
            "approval": approval,
            "checklist": checklist,
            "revision": revision,
            "fmeeting": fmeeting,
            "fmeeting_ongoing": fmeeting_ongoing,
            "fmeeting_completed": fmeeting_completed,
            "fitem": fitem,
            "fitem2": fitem2,
            "fproposal": fproposal,
            "fentry": fentry,
            "fattendance": fattendance,
            "fattachment": fattachment,
            "schedule": schedule,
            "exception": exception,
            "rule": rule,
            "task": task,
            "label": label,
            "task_checklist": task_checklist,
            "task_checklist2": task_checklist2,
            "task_attachment": task_attachment,
            "ticket": ticket,
            "ticket_closed": ticket_closed,
            "notification": notification,
            "absence": absence,
            "change_request": change_request,
            "export": export,
            "invitation": invitation,
            "role": role,
            "role_std": role_std,
            "party": party,
            "contact": contact,
        }
        ids = {key_: str(obj.pk) for key_, obj in objects.items()}
        ids.update(
            {
                "member_email": kollege_user.email,
                "guest_email": gast_user.email,
                "unbekannt": str(uuid.uuid4()),
            }
        )
        return org, ids


def _watched_models() -> list[type[Model]]:
    return [
        model
        for model in django_apps.get_models(include_auto_created=True)
        if model._meta.app_label in WATCHED_APPS and model._meta.managed and not model._meta.proxy
    ]


def _normalize(value: Any) -> Any:
    if isinstance(value, memoryview):
        return bytes(value)
    if isinstance(value, dict | list):
        return json.dumps(value, sort_keys=True, default=str)
    return value


def _snapshot() -> Snapshot:
    """Alle Zeilen der beobachteten Tabellen, je Modell nach Primärschlüssel."""
    state: Snapshot = {}
    for model in _watched_models():
        pk = model._meta.pk
        assert pk is not None
        columns = [pk.attname] + [f.attname for f in model._meta.concrete_fields if f.attname != pk.attname]
        rows = model._base_manager.values_list(*columns)
        state[model._meta.label] = {row[0]: tuple(_normalize(v) for v in row) for row in rows}
    return state


def _changes(before: Snapshot, foreign_refs: frozenset[str] = frozenset()) -> list[str]:
    """
    Geänderte Modelle gegenüber dem Ausgangsstand: ``label (+neu −gelöscht ~geändert)``.

    Enthält eine neue oder geänderte Zeile neu eine ID aus ``foreign_refs``, wird das als Verweis auf Org A vermerkt.
    """
    changed = []
    for label, rows in _snapshot().items():
        old = before.get(label, {})
        added = rows.keys() - old.keys()
        removed = old.keys() - rows.keys()
        modified = [pk for pk in rows.keys() & old.keys() if rows[pk] != old[pk]]
        if not (added or removed or modified):
            continue
        refers = any(
            str(value) in foreign_refs and value not in old.get(pk, ())
            for pk in (*added, *modified)
            for value in rows[pk]
        )
        note = ", Verweis auf Org A" if refers else ""
        changed.append(f"{label} (+{len(added)} −{len(removed)} ~{len(modified)}{note})")
    return changed


@pytest.fixture(scope="module")
def world(django_db_setup: None, django_db_blocker: Any, tmp_path_factory: pytest.TempPathFactory) -> Iterator[World]:
    """Modulweite Testdaten; Uploads landen in einem temporären MEDIA_ROOT."""
    media_root = tmp_path_factory.mktemp("work-matrix-media")
    with override_settings(MEDIA_ROOT=str(media_root)), django_db_blocker.unblock():
        created_permissions: list[str] = []
        users: list[User] = []
        source = OParlSource.objects.create(name="Matrix-RIS", url="https://ris.example.org/matrix/system")
        builder = _Builder(source, created_permissions, users)
        org_a, ids_a = builder.organization("A")
        org_b, ids_b = builder.organization("B")

        # Plattformweite Objekte: Wissensdatenbank, Partei
        category = KnowledgeBaseCategory.objects.create(name="Matrix-Hilfe", slug="matrix-hilfe")
        article = KnowledgeBaseArticle.objects.create(
            category=category, title="Artikel", slug="matrix-artikel", content="Text", is_published=True
        )
        party_group = PartyGroup.objects.create(name="Matrix-Partei", slug="matrix-partei")
        for ids in (ids_a, ids_b):
            ids["kb_article"] = str(article.pk)
            ids["party_group"] = str(party_group.pk)
        # Einstellungs-Singletons vorab anlegen: Ihr erstes Lesen legt sie an und zählte sonst als Änderung
        new_singletons = [model for model in (SiteSettings, AISettings) if not model.objects.exists()]
        for singleton in (SiteSettings, AISettings):
            cast(Any, singleton).get_settings()

        # Mitglied ohne Rechte und Basis-Mitglied (nur dashboard.view) in Org A, Administrator von Org B
        ohne_rechte = _membership(
            builder.user("ohne-rechte@matrix.example.org"), org_a, [builder.role(org_a, "Matrix ohne Rechte", [])]
        )
        basis_role = Role.objects.get(organization=org_a, name="Matrix Basis")
        basis = _membership(builder.user("basis@matrix.example.org"), org_a, [basis_role])
        admin_b = Membership.objects.get(pk=ids_b["admin"])

        clients = {
            "ohne_rechte": _login(ohne_rechte.user),
            "basis": _login(basis.user),
            "admin_b": _login(admin_b.user),
        }
        global_keys = {"kb_article", "party_group", "unbekannt"}
        org_a_refs = frozenset({str(org_a.pk), *(value for key, value in ids_a.items() if key not in global_keys)})
        built = World(ids={"A": ids_a, "B": ids_b}, clients=clients, org_a_refs=org_a_refs)
        built.baseline = _snapshot()
        yield built

        Organization.objects.filter(pk__in=[org_a.pk, org_b.pk]).delete()
        User.objects.filter(pk__in=[user.pk for user in users]).delete()
        source.delete()
        category.delete()
        party_group.delete()
        Permission.objects.filter(codename__in=created_permissions).delete()
        for model in new_singletons:
            model.objects.all().delete()


# =============================================================================
# Aufruf und Prüfung
# =============================================================================


def _resolve(value: Any, mapping: dict[str, str]) -> Any:
    if isinstance(value, list):
        return [_resolve(item, mapping) for item in value]
    if isinstance(value, str):
        return value.format_map(mapping)
    return value


def _payload(world: World, template: dict[str, Any], own: str, sub: str | None = None) -> dict[str, Any]:
    """
    Platzhalter auflösen: ``{x}`` aus der Org ``own``, ``{a_x}`` immer aus Org A.

    Mit ``sub`` stammen Unterobjekte der Pfad-Objekte (SUB_OBJECTS) aus dieser Organisation.
    """
    mapping = {**world.ids[own], **{f"a_{key}": value for key, value in world.ids["A"].items()}}
    if sub is not None:
        mapping.update({key: world.ids[sub][key] for key in SUB_OBJECTS})
    return {key: _resolve(value, mapping) for key, value in template.items()}


def _url(world: World, case: Case, slug_org: str, path_org: str) -> str:
    kwargs = {"org_slug": SLUGS[slug_org]}
    for param, key in case.path.items():
        kwargs[param] = key[1:] if key.startswith("=") else world.ids[path_org][key]
    return reverse(f"work:{case.url_name}", kwargs=kwargs)


def _send(client: Client, case: Case, url: str, data: dict[str, Any]) -> HttpResponse:
    # Cookies sichern: Die Sitzung stammt aus dem Aufbau; ein neuer Sitzungsschlüssel aus einem Aufruf
    # verschwindet mit dem Rollback und würde die folgenden Fälle abmelden.
    cookies = copy.deepcopy(client.cookies)
    try:
        if case.method == "delete":
            response = client.delete(url)
        elif case.method == "json":
            response = client.post(url, data=json.dumps(data), content_type="application/json")
        else:
            payload = dict(data)
            for field_name, filename in case.files.items():
                payload[field_name] = _upload(filename)
            response = client.post(url, payload)
    finally:
        client.cookies = cookies
    return cast(HttpResponse, response)


def _check(
    world: World,
    role: str,
    case: Case,
    *,
    slug_org: str,
    path_org: str,
    data: dict[str, Any],
    allowed: tuple[int, ...],
) -> None:
    url = _url(world, case, slug_org, path_org)
    response = _send(world.clients[role], case, url, data)
    problems = []
    if response.status_code not in allowed:
        problems.append(f"Status {response.status_code} statt {'/'.join(str(code) for code in allowed)}")
    # Unter dem Slug von Org B ist jeder neue Verweis auf ein Objekt von Org A eine Grenzverletzung
    changed = _changes(world.baseline, world.org_a_refs if slug_org == "B" else frozenset())
    if changed:
        problems.append("Datenbank geändert: " + ", ".join(changed))
    assert not problems, f"{case.method.upper()} {url} [{case.name}, {role}]: " + "; ".join(problems)


def _grenzverletzungen(before: Snapshot, refs: frozenset[str]) -> list[str]:
    """Geänderte oder gelöschte Zeilen von Org A sowie neue Verweise auf Org A gegenüber dem Ausgangsstand."""

    def gehoert_zu_a(pk: Any, values: tuple[Any, ...]) -> bool:
        return str(pk) in refs or any(str(value) in refs for value in values)

    found = []
    for label, rows in _snapshot().items():
        old = before.get(label, {})
        for pk in old.keys() - rows.keys():
            if gehoert_zu_a(pk, old[pk]):
                found.append(f"{label}: Zeile von Org A gelöscht")
        for pk in rows:
            if pk in old and rows[pk] != old[pk] and gehoert_zu_a(pk, old[pk]):
                found.append(f"{label}: Zeile von Org A geändert")
            elif rows[pk] != old.get(pk) and any(
                str(value) in refs and value not in old.get(pk, ()) for value in rows[pk]
            ):
                found.append(f"{label}: neuer Verweis auf Org A")
    return found


def _params(cases: list[Case]) -> list[Any]:
    return [pytest.param(case, id=case.name) for case in cases]


# =============================================================================
# Matrix
# =============================================================================


@pytest.mark.parametrize("case", _params([c for c in CASES if not c.self_service]))
def test_mitglied_ohne_rechte(world: World, case: Case) -> None:
    """Mitglied von Org A ohne Berechtigungen: 403/404, keine Änderung."""
    _check(
        world,
        "ohne_rechte",
        case,
        slug_org="A",
        path_org="A",
        data=_payload(world, case.data, "A"),
        allowed=(403, 404),
    )


@pytest.mark.parametrize("case", _params(CASES))
def test_fremde_organisation(world: World, case: Case) -> None:
    """Administrator von Org B unter dem Slug von Org A: 403, keine Änderung."""
    _check(world, "admin_b", case, slug_org="A", path_org="A", data=_payload(world, case.data, "A"), allowed=(403,))


@pytest.mark.parametrize("case", _params([c for c in CASES if c.idor and c.has_path_objects]))
def test_fremde_objekte_im_pfad(world: World, case: Case) -> None:
    """Administrator von Org B unter eigenem Slug mit Objekt-IDs aus Org A im Pfad: 403/404, keine Änderung."""
    _check(
        world,
        "admin_b",
        case,
        slug_org="B",
        path_org="A",
        data=_payload(world, case.data, "B", sub="A"),
        allowed=(403, 404),
    )


@pytest.mark.parametrize(
    ("case", "variant"),
    [pytest.param(case, variant, id=f"{case.name}:{variant}") for case in CASES for variant in case.foreign],
)
def test_fremde_objekte_in_nutzdaten(world: World, case: Case, variant: str) -> None:
    """Administrator von Org B mit eigenen Pfad-Objekten, aber IDs aus Org A in den Nutzdaten: Org A bleibt unberührt."""
    data = _payload(world, {**case.data, **case.foreign[variant]}, "B")
    url = _url(world, case, "B", "B")
    response = _send(world.clients["admin_b"], case, url, data)
    problems = _grenzverletzungen(world.baseline, world.org_a_refs)
    if response.status_code >= 500:
        problems.append(f"Status {response.status_code}")
    assert not problems, f"{case.method.upper()} {url} [{case.name}:{variant}, admin_b]: " + "; ".join(problems)


@pytest.mark.parametrize("case", _params([c for c in CASES if c.same_org]))
def test_selbstbedienung_fremdes_objekt(world: World, case: Case) -> None:
    """Basis-Mitglied von Org A mit dem Objekt einer Kollegin: 403/404, keine Änderung."""
    _check(world, "basis", case, slug_org="A", path_org="A", data=_payload(world, case.data, "A"), allowed=(403, 404))


# =============================================================================
# Vollständigkeit
# =============================================================================


def test_matrix_deckt_alle_schreibenden_endpunkte() -> None:
    """Jede Work-URL, deren View POST/PUT/PATCH/DELETE annimmt, steht in der Matrix oder in NICHT_IN_MATRIX."""
    covered = {case.url_name for case in CASES}
    missing = []
    for pattern in work_urls.urlpatterns:
        name = getattr(pattern, "name", None)
        callback = getattr(pattern, "callback", None)
        view_class = getattr(callback, "view_class", None)
        if not name or view_class is None or name in NICHT_IN_MATRIX:
            continue
        allowed = getattr(callback, "view_initkwargs", {}).get("http_method_names", view_class.http_method_names)
        writes = any(method in allowed and hasattr(view_class, method) for method in ("post", "put", "patch", "delete"))
        if writes and name not in covered:
            missing.append(name)
    assert not missing, f"Schreibende Endpunkte ohne Fall in der Matrix: {missing}"
    assert not set(SELBSTBEDIENUNG) - covered, "SELBSTBEDIENUNG nennt URL-Namen ohne Fall"
