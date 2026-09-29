# SPDX-License-Identifier: AGPL-3.0-or-later
"""
Gliederung beim Dokument-Import (Issue #620).

Die Eingaben stammen aus echten Anträgen (Wortlaut gekürzt): pypdf lieferte je nach Erzeuger
eine ganze Seite als eine Zeile, jedes Wort in einer eigenen Zeile oder – im Layout-Modus –
zerrissene Wörter bei Blocksatz. Im Editor landete früher eine ganze Seite in einem Absatz.
"""

from __future__ import annotations

import io

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from apps.work.motions import import_text


class TestLayoutMitWortabstaenden:
    def test_zerrissene_woerter_aus_dem_layout_werden_repariert(self) -> None:
        """Änderungsantrag mit Zeichenabständen: Layout zerreißt Wörter, der einfache Modus hat keine Zeilen."""
        einfach = (
            "   Münster, 21.04.2026 Änderungsantrag V/0106/2026 z. B.: • Stopp-Regelung an der Einfahrt "
            "• Temporeduzierung  • Farbliche Markierungen  Begründung: Die Planung ist zeitaufwendig."
        )
        layout = (
            "                                                   Münster, 21.04.2026\n\n"
            "Änderungsantrag\n"
            "V/   0   1   0   6   /   2   0   2   6\n\n"
            "      •      Stopp-Regelung an der Einfahrt\n\n"
            "      •      Te      m      p      o      r     e      d      u      z      i      e      r      u      n      g\n\n"
            "      •      Fa r b l i c h e   M a r k i e r u n ge n\n\n\n"
            "Begründung:\n\n"
            "Die       Planung       ist       zeitaufwendig."
        )
        zeilen = import_text.merge_layout_lines(einfach, layout)
        assert "• Temporeduzierung" in zeilen
        assert "• Farbliche Markierungen" in zeilen
        assert "V/0106/2026" in zeilen
        assert "Die Planung ist zeitaufwendig." in zeilen
        assert "" in zeilen  # Leerzeilen tragen die Absatzabstände

    def test_zeile_ausserhalb_des_textflusses_bleibt_mit_einfachen_abstaenden(self) -> None:
        # Wasserzeichen mitten in der Zeile: im einfachen Text steht es an anderer Stelle
        zeilen = import_text.merge_layout_lines(
            "ENTWURF Die Stadt Münster verfolgt", "Die   Stadt   MünsterENTWURF verfolgt"
        )
        assert zeilen == ["Die Stadt MünsterENTWURF verfolgt"]


class TestAbsaetze:
    def test_umbrochene_zeilen_werden_ein_absatz_nummern_eigene_absaetze(self) -> None:
        seite = [
            "Ratsantrag",
            "Münster, 28.09.2026",
            "",
            "Der Rat möge beschließen:",
            "1. Der Rat bekräftigt, dass digitale Gremienarbeit in Münster Standard ist und bleibt. Bera-",
            "tungs- und Arbeitsunterlagen sind grundsätzlich digital über das Ratsinformationssystem",
            "bereitzustellen.",
            "2. Die Verwaltung wird beauftragt, die sog. Ratspost einzustellen. Die Umstellung ist spätes-",
            "tens zum 01.01.2027 vorzunehmen.",
            "",
            "Begründung:",
            "Die Stadt Münster verfolgt den Anspruch, die Digitalisierung konsequent und glaubwürdig um-",
            "zusetzen. Dazu gehört, dass nicht nur Verwaltungsprozesse modernisiert werden, sondern die",
            "politische Arbeit selbst.",
            "Seite 1 von 2",
        ]
        absaetze = import_text.paragraphs_from_lines([seite])
        assert absaetze == [
            "Ratsantrag",
            "Münster, 28.09.2026",
            "Der Rat möge beschließen:",
            "1. Der Rat bekräftigt, dass digitale Gremienarbeit in Münster Standard ist und bleibt. "
            "Beratungs- und Arbeitsunterlagen sind grundsätzlich digital über das Ratsinformationssystem "
            "bereitzustellen.",
            "2. Die Verwaltung wird beauftragt, die sog. Ratspost einzustellen. Die Umstellung ist spätestens "
            "zum 01.01.2027 vorzunehmen.",
            "Begründung:",
            "Die Stadt Münster verfolgt den Anspruch, die Digitalisierung konsequent und glaubwürdig umzusetzen. "
            "Dazu gehört, dass nicht nur Verwaltungsprozesse modernisiert werden, sondern die politische Arbeit "
            "selbst.",
        ]

    def test_bindestrich_woerter_und_ergaenzungsstrich_bleiben(self) -> None:
        zeilen = [
            "Die Verwaltung wird beauftragt, weitere Ad-hoc-Maßnahmen und eine gemeinsame Plattform für die Rats-",
            "und Gremienarbeit zu prüfen, auf der Vereine und Institutionen per E-",
            "Mail erreichbar sind.",
        ]
        (absatz,) = import_text.paragraphs_from_lines([zeilen])
        assert "Rats- und Gremienarbeit" in absatz
        assert "per E-Mail erreichbar" in absatz

    def test_absatz_laeuft_ueber_die_seitengrenze(self) -> None:
        voll = "Die Verwaltung wird beauftragt, die durch die Umstellung tatsächlich erzielten Einsparungen"
        seite1 = [voll, voll, "und Mehrkosten transparent und getrennt nach Kostenarten darzustellen, damit der Rat"]
        seite2 = ["3", "prüfen kann, ob die Umstellung wirtschaftlich ist.", "", "Begründung:"]
        absaetze = import_text.paragraphs_from_lines([seite1, seite2])
        assert absaetze[0].endswith("damit der Rat prüfen kann, ob die Umstellung wirtschaftlich ist.")
        assert absaetze[1] == "Begründung:"

    def test_anderthalbzeiliger_abstand_trennt_nicht_jede_zeile(self) -> None:
        """Layout-Modus setzt bei weitem Zeilenabstand zwischen jede Zeile eine Leerzeile."""
        lang = "Die Verwaltung wird gebeten zu prüfen, ob auf der Homepage der Stadt wieder die einzelnen"
        zeilen = [lang, "", lang, "", lang, "", lang, "", "Ortsteile aufgeführt werden können.", "", "", "Begründung:"]
        zeilen += ["", lang, "", "Vereine und Institutionen haben eine Bringschuld."]
        absaetze = import_text.paragraphs_from_lines([zeilen])
        assert len(absaetze) == 3, absaetze
        assert absaetze[0].endswith("Ortsteile aufgeführt werden können.")
        assert absaetze[1] == "Begründung:"

    def test_texterkennung_mit_leerzeilen(self) -> None:
        text = "Antrag der Fraktion\n\nDer Rat beschließt, 50 Bäume\nzu pflanzen.\n\n1. Standorte prüfen"
        assert import_text.text_paragraphs(text) == [
            "Antrag der Fraktion",
            "Der Rat beschließt, 50 Bäume zu pflanzen.",
            "1. Standorte prüfen",
        ]

    def test_html_maskiert_und_loest_ligaturen_auf(self) -> None:
        html = import_text.paragraphs_to_html(["Einpﬂege <sofort> & ﬁnal"])
        assert html == "<p>Einpflege &lt;sofort&gt; &amp; final</p>"


def _nummerierung(dokument: docx.document.Document) -> int:
    """Nummerierung wie in Word: Ebene 0 „1.“, Ebene 1 „a)“."""
    numbering = dokument.part.numbering_part.element
    numbering.append(
        parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90">'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
            '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%2)"/></w:lvl>'
            "</w:abstractNum>"
        )
    )
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
    return 90


def _punkt(dokument: docx.document.Document, text: str, num_id: int, ebene: int = 0) -> None:
    absatz = dokument.add_paragraph(text, style="List Paragraph")
    num_pr = absatz._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_numId().val = num_id
    num_pr.get_or_add_ilvl().val = ebene


def _antrag_docx() -> bytes:
    """Aufbau wie ein echter Ratsantrag aus Word (Liste, Unterpunkte, Unterschriften-Tabelle)."""
    dokument = docx.Document()
    num_id = _nummerierung(dokument)
    leer = dokument.add_table(rows=1, cols=3)  # Briefkopf-Raster ohne Text
    assert leer is not None
    dokument.add_paragraph("Ratsantrag\tMünster, 28.09.2026")
    dokument.add_heading("Antrag zur Stärkung der digitalen Ratsarbeit", 1)
    dokument.add_paragraph("Der Rat möge beschließen:", style="List Paragraph")  # ohne Nummerierung
    _punkt(dokument, "Der Rat bekräftigt die digitale Gremienarbeit.", num_id)
    _punkt(dokument, "Die Strategie verfolgt folgende Ziele:", num_id)
    _punkt(dokument, "Transparenz", num_id, 1)
    _punkt(dokument, "Datenschutz", num_id, 1)
    _punkt(dokument, "Die Verwaltung stellt die Ratspost ein.", num_id)
    begruendung = dokument.add_paragraph()
    begruendung.add_run("Begründung:").bold = True
    begruendung.add_run().add_break()
    begruendung.add_run("Die Stadt verfolgt den Anspruch, ")
    # Link (Läufe in w:hyperlink fehlten früher ganz)
    rel_id = dokument.part.relate_to("https://www.example.org/ris", RT.HYPERLINK, is_external=True)
    begruendung._p.append(
        parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{rel_id}"><w:r><w:t>das RIS</w:t></w:r></w:hyperlink>')
    )
    begruendung.add_run(" zu stärken.")
    # Überarbeitung: gelöschter Text entfällt
    begruendung._p.append(
        parse_xml(f'<w:del {nsdecls("w")} w:id="1" w:author="x"><w:r><w:delText>GELÖSCHT</w:delText></w:r></w:del>')
    )
    dokument.add_paragraph("gez.")
    tabelle = dokument.add_table(rows=3, cols=2)
    tabelle.cell(0, 0).text = "Albert Wenzel"
    tabelle.cell(0, 1).text = "Noah Börnhorst"
    tabelle.cell(2, 0).text = "Katja Martinewski"
    puffer = io.BytesIO()
    dokument.save(puffer)
    return puffer.getvalue()


class TestDocx:
    def test_liste_unterpunkte_tabelle_und_umbrueche(self) -> None:
        html = import_text.docx_to_html(_antrag_docx())

        assert "<p>Der Rat möge beschließen:</p>" in html  # Listenabsatz ohne Nummer bleibt Absatz
        assert "<h1>Antrag zur Stärkung der digitalen Ratsarbeit</h1>" in html
        assert "<p>Ratsantrag Münster, 28.09.2026</p>" in html
        assert (
            "<ol><li><p>Der Rat bekräftigt die digitale Gremienarbeit.</p></li>"
            "<li><p>Die Strategie verfolgt folgende Ziele:</p>"
            '<ol type="a"><li><p>Transparenz</p></li><li><p>Datenschutz</p></li></ol></li>'
            "<li><p>Die Verwaltung stellt die Ratspost ein.</p></li></ol>"
        ) in html
        assert "<strong>Begründung:</strong><br>Die Stadt verfolgt den Anspruch, " in html
        assert '<a href="https://www.example.org/ris">das RIS</a> zu stärken.' in html
        assert "GELÖSCHT" not in html
        # Unterschriften-Tabelle bleibt, leere Zeile und leeres Raster entfallen
        assert html.count("<table>") == 1
        assert "<td><p>Albert Wenzel</p></td><td><p>Noah Börnhorst</p></td>" in html
        assert "Katja Martinewski" in html
        assert html.count("<tr>") == 2

    def test_nummerierung_laeuft_nach_zwischenabsatz_weiter(self) -> None:
        dokument = docx.Document()
        num_id = _nummerierung(dokument)
        _punkt(dokument, "Erstens", num_id)
        _punkt(dokument, "Zweitens", num_id)
        dokument.add_paragraph("Zwischentext")
        _punkt(dokument, "Drittens", num_id)
        puffer = io.BytesIO()
        dokument.save(puffer)
        html = import_text.docx_to_html(puffer.getvalue())
        assert '<ol start="3"><li><p>Drittens</p></li></ol>' in html

    def test_suchtext_aus_html(self) -> None:
        text = import_text.html_to_text("<h1>Titel</h1><p>A &amp; B<br>C</p><table><tr><td><p>X</p></td></tr></table>")
        assert text == "Titel\nA & B\nC\nX"
