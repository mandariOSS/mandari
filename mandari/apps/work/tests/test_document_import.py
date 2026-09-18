# SPDX-License-Identifier: AGPL-3.0-or-later
"""
Dokument-Import (PDF/DOCX) in Work: Die Seite war ohne Funktion, weil ihr Skript in einem
nicht existierenden Template-Block stand; zudem gingen Fehlschläge verloren, wenn genau eine
Datei gelang (direkter Sprung in den Editor vor den Fehlermeldungen).
"""

from __future__ import annotations

import io
from typing import Any

import docx
import pytest
from django.contrib.messages import get_messages
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.work.motions.models import Motion

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx() -> bytes:
    dokument = docx.Document()
    dokument.add_heading("Antrag: Mehr Bäume in der Innenstadt", 1)
    dokument.add_paragraph("Der Rat beschließt, 50 Bäume zu pflanzen.")
    puffer = io.BytesIO()
    dokument.save(puffer)
    return puffer.getvalue()


@pytest.fixture
def mitglied(org: Any, make_member: Any) -> Any:
    return make_member(org, ["dashboard.view", "motions.view", "motions.create", "motions.edit"])


@pytest.mark.django_db
def test_import_seite_liefert_ihr_skript(org: Any, mitglied: Any, client_for: Any) -> None:
    seite = client_for(mitglied.user).get(f"/work/{org.slug}/documents/import/").content.decode()
    assert 'x-data="importForm()"' in seite and "function importForm()" in seite


@pytest.mark.django_db
def test_docx_wird_bearbeitbares_dokument(org: Any, mitglied: Any, client_for: Any) -> None:
    datei = SimpleUploadedFile("baeume.docx", _docx(), content_type=DOCX)
    antwort = client_for(mitglied.user).post(f"/work/{org.slug}/documents/import/", {"import_files": [datei]})
    assert antwort.status_code == 302 and "/documents/" in antwort["Location"]
    motion = Motion.objects.get(organization=org)
    assert "50 Bäume" in (motion.content or "")


@pytest.mark.django_db
def test_fehlschlag_wird_auch_neben_erfolg_gemeldet(org: Any, mitglied: Any, client_for: Any) -> None:
    gut = SimpleUploadedFile("baeume.docx", _docx(), content_type=DOCX)
    kaputt = SimpleUploadedFile("kaputt.docx", b"PK\x03\x04 kein Word", content_type=DOCX)
    antwort = client_for(mitglied.user).post(f"/work/{org.slug}/documents/import/", {"import_files": [gut, kaputt]})
    texte = [str(m) for m in get_messages(antwort.wsgi_request)]
    assert any("erfolgreich importiert" in t for t in texte)
    assert any("fehlgeschlagen" in t or "übersprungen" in t for t in texte), texte


@pytest.mark.django_db
def test_ohne_recht_kein_import(org: Any, make_member: Any, client_for: Any) -> None:
    lesend = make_member(org, ["dashboard.view", "motions.view"])
    datei = SimpleUploadedFile("baeume.docx", _docx(), content_type=DOCX)
    antwort = client_for(lesend.user).post(f"/work/{org.slug}/documents/import/", {"import_files": [datei]})
    assert antwort.status_code in (302, 403)
    assert not Motion.objects.filter(organization=org).exists()
