# SPDX-License-Identifier: AGPL-3.0-or-later
"""
Sichern vor jedem Schritt, der die Seite verlässt oder neu lädt (Issue #1018), im neuen Antragseditor.

Mit Live-Zusammenarbeit (ASGI-Testserver mit In-Memory-Channel-Layer, wie tests_e2e/test_editor_kollaboration.py):

- Titel, Dokumenttyp und Briefpapier laufen nicht über Yjs. Sie werden per POST gesichert, ohne den Text zu schicken,
  und der Editor gilt erst danach als gespeichert – eine Bestätigung der Verbindung (yjs_saved) allein reicht nicht.
- „Freigeben“ sichert vorher den gemeinsamen Stand. Das auslösende Fenster bekommt keine Aufforderung zum Neuladen,
  das andere schon; die letzten Eingaben stehen in der gesperrten Fassung.
- „Herunterladen als Word“ sichert vorher und lädt die Datei, ohne die Seite zu verlassen.
- „Vorschau (PDF)“, Links auf andere Seiten und der Teilen-Dialog sichern vorher und gehen ohne Rückfrage.

Ohne Verbindung (WSGI-Live-Server, Solo-Modus): Federführung und Zustimmen aus der Seitenleiste bzw. Infozeile sichern
vorher den Text und laden ohne Rückfrage „Website verlassen?“ neu.

Das regelmäßige Sichern der Verbindung (alle 10 s) ist in diesen Tests abgeschaltet: Sonst hinge es vom Zufall ab, ob
die letzten Eingaben schon gespeichert sind. Gesichert wird nur, was die geprüften Schritte selbst sichern.
"""

from __future__ import annotations

import json
from collections.abc import Callable
from pathlib import Path
from typing import Any, cast

import pytest

from apps.work.motions.models import Motion, MotionApproval, MotionType, OrganizationLetterhead
from tests_e2e.conftest import login_via_form, wait_for_bundle
from tests_e2e.test_editor import ALPINE_EDITOR, PASSWORD, PERMISSIONS, PROSEMIRROR
from tests_e2e.test_editor_kollaboration import (
    decrypted_content,
    editor_oeffnen,
    editor_url,
    tippen,
    warte_auf_editor_text,
)

playwright_sync = pytest.importorskip("playwright.sync_api")
expect = playwright_sync.expect

pytestmark = pytest.mark.django_db(transaction=True)

DESKTOP = {"width": 1440, "height": 900}
#: Das regelmäßige Sichern des Kollaborationsanbieters (setInterval mit 10 s) auf eine Stunde strecken
OHNE_REGELMAESSIGES_SICHERN = """
(() => {
    const original = window.setInterval
    window.setInterval = (fn, ms, ...rest) => original(fn, ms === 10000 ? 3600000 : ms, ...rest)
})()
"""
GESPEICHERT_ANZEIGE = ".ke-gespeichert"


@pytest.fixture(autouse=True)
def neues_design(db: Any, org: Any) -> None:
    org.work_new_design = True
    org.save(update_fields=["work_new_design"])


def _mitglied(make_member: Any, org: Any, email: str, permissions: list[str]) -> Any:
    membership = make_member(org, permissions, email=email)
    membership.user.set_password(PASSWORD)
    membership.user.save(update_fields=["password"])
    return membership


@pytest.fixture
def anna(db: Any, org: Any, make_member: Any) -> Any:
    """Autorin: bearbeitet, teilt und steuert den Ablauf ihres Dokuments."""
    return _mitglied(make_member, org, "anna-sichern@example.org", [*PERMISSIONS, "motions.share"])


@pytest.fixture
def bernd(db: Any, org: Any, make_member: Any) -> Any:
    """Zweites Mitglied, darf alle Dokumente bearbeiten (Kollaboration in fremden Dokumenten)."""
    return _mitglied(make_member, org, "bernd-sichern@example.org", [*PERMISSIONS, "motions.edit_all"])


@pytest.fixture
def dokument(org: Any, anna: Any) -> Motion:
    motion = Motion.objects.create(
        organization=org, author=anna, title="Antrag Radweg", summary="Kurzfassung", visibility="organization"
    )
    cast(Any, motion).set_content_encrypted("<p>Alt</p>")
    motion.save()
    return motion


class WsNachrichten:
    """JSON-Nachrichten der Kollaborationsverbindung einer Seite (auch über ein Neuladen hinweg)."""

    def __init__(self, page: Any) -> None:
        self.empfangen: list[dict[str, Any]] = []
        self.gesendet: list[dict[str, Any]] = []
        #: Gemeinsame Liste mit anderen Ereignissen der Seite: Typen empfangener Nachrichten in der Reihenfolge
        self.protokoll: list[str] | None = None
        page.on("websocket", self._verbindung)

    def _verbindung(self, ws: Any) -> None:
        if "/ws/documents/" not in ws.url:
            return
        ws.on("framereceived", lambda payload: self._merken(self.empfangen, payload, self.protokoll))
        ws.on("framesent", lambda payload: self._merken(self.gesendet, payload, None))

    @staticmethod
    def _merken(liste: list[dict[str, Any]], payload: Any, protokoll: list[str] | None) -> None:
        if isinstance(payload, bytes):
            payload = payload.decode("utf-8", "replace")
        try:
            nachricht = json.loads(payload)
        except ValueError:
            return
        if isinstance(nachricht, dict):
            liste.append(nachricht)
            if protokoll is not None:
                protokoll.append(str(nachricht.get("type")))

    def empfangene_typen(self) -> list[str]:
        return [str(n.get("type")) for n in self.empfangen]


def dialoge_merken(page: Any) -> list[str]:
    """Rückfragen des Browsers (beforeunload, confirm) merken und ablehnen: Ein Schritt mit Rückfrage gilt als Fehler."""
    dialoge: list[str] = []

    def ablehnen(dialog: Any) -> None:
        dialoge.append(dialog.type)
        dialog.dismiss()

    page.on("dialog", ablehnen)
    return dialoge


def vorbereiten(page: Any) -> None:
    page.set_viewport_size(DESKTOP)
    page.add_init_script(OHNE_REGELMAESSIGES_SICHERN)


def kein_neuladen_markieren(page: Any) -> None:
    page.evaluate("() => { window.__ohneNeuladen = true }")


def nicht_neu_geladen(page: Any) -> bool:
    return bool(page.evaluate("() => window.__ohneNeuladen === true"))


def warten_bis(page: Any, bedingung: Callable[[], bool], beschreibung: str, sekunden: float = 15) -> None:
    for _ in range(int(sekunden * 10)):
        if bedingung():
            return
        page.wait_for_timeout(100)
    raise AssertionError(f"Nicht eingetreten: {beschreibung}")


def datei_menue(page: Any, punkt: str) -> Any:
    page.locator(".ke-kopf [data-menue-knopf=datei]").click()
    return page.locator(".ke-kopf").get_by_role("menuitem", name=punkt)


def docx_text(pfad: Path) -> str:
    from docx import Document

    dokument = Document(str(pfad))
    teile = [absatz.text for absatz in dokument.paragraphs]
    for tabelle in dokument.tables:
        teile.extend(zelle.text for zeile in tabelle.rows for zelle in zeile.cells)
    return "\n".join(teile)


class TestMitVerbindung:
    def test_titel_typ_und_briefpapier_gesichert(
        self, page: Any, asgi_server: Any, org: Any, anna: Any, dokument: Motion
    ) -> None:
        typ = MotionType.objects.create(organization=org, name="Anfrage", slug="anfrage")
        briefkopf = OrganizationLetterhead.objects.create(organization=org, name="Fraktion", kind="generated")
        vorbereiten(page)
        ws = WsNachrichten(page)
        login_via_form(page, asgi_server.url, anna.user.email, PASSWORD)
        url = editor_url(asgi_server, org, dokument)
        editor_oeffnen(page, url, "Alt")
        kein_neuladen_markieren(page)

        # POST-Anfragen an den Editor zurückhalten, um den Zwischenstand zu prüfen
        gehalten: list[Any] = []
        freigegeben: list[bool] = []

        def halten(route: Any) -> None:
            if route.request.method == "POST" and not freigegeben:
                gehalten.append(route)
            else:
                route.continue_()

        page.route(f"**/documents/{dokument.id}/", halten)

        page.fill("#ke-titel", "Antrag Radweg Hauptstraße")
        page.locator("[data-menue-knopf=typ]").click()
        page.get_by_role("menuitemradio", name="Anfrage").click()
        page.locator("[data-menue-knopf=briefpapier]").click()
        page.get_by_role("menuitemradio", name="Fraktion").click()
        tippen(page, " Neu")

        # Die Verbindung sichert den Text (Tab kurz verborgen); das allein gilt nicht als gespeichert
        bestaetigt = ws.empfangene_typen().count("yjs_saved")
        page.evaluate(
            """() => {
                Object.defineProperty(document, 'visibilityState', { value: 'hidden', configurable: true })
                document.dispatchEvent(new Event('visibilitychange'))
                delete document.visibilityState
            }"""
        )
        warten_bis(page, lambda: ws.empfangene_typen().count("yjs_saved") > bestaetigt, "yjs_saved nach dem Sichern")
        assert page.evaluate(f"() => {ALPINE_EDITOR}.hasUnsavedChanges()") is True
        expect(page.locator(GESPEICHERT_ANZEIGE)).to_contain_text("Ungespeicherte Änderungen", use_inner_text=True)

        # Kurz danach gehen Titel, Typ und Briefpapier per POST – ohne den Text
        warten_bis(page, lambda: bool(gehalten), "POST mit Titel, Dokumenttyp und Briefpapier")
        anfrage = gehalten[0].request.post_data or ""
        assert 'name="title"' in anfrage and "Antrag Radweg Hauptstraße" in anfrage
        assert 'name="content"' not in anfrage, "mit Verbindung darf der POST den Text nicht schicken"
        freigegeben.append(True)
        for route in gehalten:
            route.continue_()

        page.wait_for_function(f"() => {ALPINE_EDITOR}.hasUnsavedChanges() === false", timeout=15000)
        expect(page.locator(GESPEICHERT_ANZEIGE)).to_contain_text("Gespeichert", use_inner_text=True)
        dokument.refresh_from_db()
        assert dokument.title == "Antrag Radweg Hauptstraße"
        assert dokument.document_type_id == typ.id
        assert dokument.letterhead_id == briefkopf.id
        assert "Alt Neu" in decrypted_content(dokument)
        assert nicht_neu_geladen(page), "das eigene Fenster darf nicht neu laden"
        assert "reload" not in ws.empfangene_typen()

    def test_freigeben_sichert_vorher_und_laedt_nur_die_anderen(
        self, new_context: Any, asgi_server: Any, org: Any, anna: Any, bernd: Any, dokument: Motion
    ) -> None:
        dokument.status = "internal_review"
        dokument.save(update_fields=["status"])
        page_a, page_b = new_context().new_page(), new_context().new_page()
        for page in (page_a, page_b):
            vorbereiten(page)
        ws_a, ws_b = WsNachrichten(page_a), WsNachrichten(page_b)
        dialoge_a = dialoge_merken(page_a)
        url = editor_url(asgi_server, org, dokument)
        login_via_form(page_a, asgi_server.url, anna.user.email, PASSWORD)
        editor_oeffnen(page_a, url, "Alt")
        login_via_form(page_b, asgi_server.url, bernd.user.email, PASSWORD)
        editor_oeffnen(page_b, url, "Alt")

        # Während der Abstimmung öffnet der Editor im Modus „Vorschlagen“: zum Tippen auf „Bearbeiten“
        page_a.evaluate(f"() => {ALPINE_EDITOR}.modusSetzen('bearbeiten')")
        tippen(page_a, " Letzter Satz.")
        warte_auf_editor_text(page_b, "Alt Letzter Satz.")
        assert "Letzter Satz." not in decrypted_content(dokument), "Vorbedingung: noch nicht gesichert"

        page_a.get_by_test_id("ablauf-info").get_by_role("button", name="Freigeben …").click()
        with page_a.expect_navigation(timeout=20000):
            page_a.get_by_test_id("dialog-freigeben").get_by_role("button", name="Freigeben", exact=True).click()
        wait_for_bundle(page_a)

        dokument.refresh_from_db()
        assert dokument.status == "approved"
        assert "Alt Letzter Satz." in decrypted_content(dokument), "die letzten Eingaben fehlen in der Fassung"
        # Keine Aufforderung „an anderer Stelle gespeichert“ an das auslösende Fenster (die Herabstufung des eigenen
        # Zugriffs durch die Sperre meldet der Server weiterhin, mit Grund)
        neuladen_a = [n for n in ws_a.empfangen if n.get("type") == "reload" and not n.get("reason")]
        assert not neuladen_a, f"das auslösende Fenster lädt selbst neu: {neuladen_a}"
        assert dialoge_a == []
        # Das andere Fenster wird zum Neuladen aufgefordert (gesperrter Status)
        warten_bis(
            page_b,
            lambda: any(n.get("type") == "reload" and not n.get("reason") for n in ws_b.empfangen),
            "Aufforderung zum Neuladen im anderen Fenster",
        )
        # Das andere Fenster lädt nach 0,8 s selbst neu; erst danach weiter
        page_b.wait_for_timeout(1500)

        # „Als eingereicht markieren“ sichert ebenfalls vorher (Bernd darf trotz Sperre bearbeiten)
        editor_oeffnen(page_b, url, "Letzter Satz.")
        dialoge_b = dialoge_merken(page_b)
        tippen(page_b, " Eingereicht.")
        page_b.get_by_test_id("ablauf-info").get_by_role("button", name="Einreichen …").click()
        with page_b.expect_navigation(timeout=20000):
            page_b.get_by_test_id("dialog-einreichen").get_by_role("button", name="Als eingereicht markieren").click()
        dokument.refresh_from_db()
        assert dokument.status == "submitted"
        assert "Alt Letzter Satz. Eingereicht." in decrypted_content(dokument)
        assert dialoge_b == []

    def test_word_herunterladen_sichert_ohne_seitenwechsel(
        self, page: Any, asgi_server: Any, org: Any, anna: Any, dokument: Motion
    ) -> None:
        vorbereiten(page)
        dialoge = dialoge_merken(page)
        login_via_form(page, asgi_server.url, anna.user.email, PASSWORD)
        url = editor_url(asgi_server, org, dokument)
        editor_oeffnen(page, url, "Alt")
        kein_neuladen_markieren(page)
        # Verlässt der Klick die Seite (Navigation zur Datei), feuert beforeunload – auch wenn daraus ein Download wird
        page.evaluate(
            "() => { window.__verlassen = 0; addEventListener('beforeunload', () => { window.__verlassen += 1 }) }"
        )
        tippen(page, " Für die Word-Datei.")

        with page.expect_download(timeout=20000) as download:
            datei_menue(page, "Herunterladen als Word").click()

        assert "Für die Word-Datei." in docx_text(Path(download.value.path()))
        assert "Für die Word-Datei." in decrypted_content(dokument)
        assert page.evaluate("() => window.__verlassen") == 0, "Herunterladen darf die Seite nicht verlassen"
        assert dialoge == [], "Herunterladen darf nicht nach dem Verlassen der Seite fragen"
        assert nicht_neu_geladen(page) and page.url == url

    def test_vorschau_links_und_teilen_sichern_vorher(
        self, page: Any, asgi_server: Any, org: Any, anna: Any, dokument: Motion
    ) -> None:
        vorbereiten(page)
        dialoge = dialoge_merken(page)
        login_via_form(page, asgi_server.url, anna.user.email, PASSWORD)
        url = editor_url(asgi_server, org, dokument)
        editor_oeffnen(page, url, "Alt")

        # „Vorschau (PDF)“: neue Registerkarte mit dem gesicherten Stand, im Browser statt als Download
        tippen(page, " Für die Vorschau.")
        with page.context.expect_event("response", lambda r: "vorschau=1" in r.url, timeout=30000) as antwort:
            datei_menue(page, "Vorschau (PDF)").click()
        assert antwort.value.headers.get("content-disposition", "").startswith("inline")
        assert "Für die Vorschau." in decrypted_content(dokument)
        for andere in page.context.pages:
            if andere is not page:
                andere.close()

        # Link im Kopf: sichern, dann ohne Rückfrage gehen
        tippen(page, " Vor dem Gehen.")
        with page.expect_navigation(timeout=20000):
            page.locator(".ke-kopf a[aria-label='Zu den Dokumenten']").click()
        assert "Vor dem Gehen." in decrypted_content(dokument)

        # Teilen-Dialog: sichert vor dem Absenden, danach lädt die Seite ohne Rückfrage neu (HX-Refresh)
        editor_oeffnen(page, url, "Vor dem Gehen.")
        tippen(page, " Vor dem Teilen.")
        datei_menue(page, "Teilen …").click()
        teilen = page.get_by_role("dialog", name="Dokument teilen")
        expect(teilen).to_be_visible()
        with page.expect_navigation(timeout=20000):
            teilen.get_by_role("button", name="Speichern").click()
        assert "Vor dem Teilen." in decrypted_content(dokument)
        assert dialoge == []


class TestOhneVerbindung:
    """Seitenleiste und Infozeile im Solo-Modus: Ohne Sichern gingen die Eingaben beim Neuladen verloren."""

    def test_federfuehrung_und_zustimmen_sichern_vorher(
        self, page: Any, live_server: Any, org: Any, anna: Any, bernd: Any, dokument: Motion
    ) -> None:
        page.set_viewport_size(DESKTOP)
        dialoge = dialoge_merken(page)
        login_via_form(page, live_server.url, anna.user.email, PASSWORD)
        url = f"{live_server.url}/work/{org.slug}/documents/{dokument.id}/"
        page.goto(url)
        wait_for_bundle(page)
        page.wait_for_selector(PROSEMIRROR, timeout=15000)
        page.wait_for_function(f"() => {ALPINE_EDITOR}.collabEnabled === false", timeout=15000)

        # Federführung (Details): sichert den Text, lädt ohne Rückfrage neu
        tippen(page, " Vor der Federführung.")
        datei_menue(page, "Details und Zuständigkeit").click()
        with page.expect_navigation(timeout=20000):
            page.locator(".ke-panel select[name=responsible]").select_option(str(bernd.id))
        dokument.refresh_from_db()
        assert dokument.responsible_id == bernd.id
        assert "Alt Vor der Federführung." in decrypted_content(dokument)

        # Zustimmen in der Infozeile (Abstimmung läuft, Anna ist angefragt)
        dokument.status = "internal_review"
        dokument.save(update_fields=["status"])
        anfrage = MotionApproval.objects.create(motion=dokument, approver=anna, approval_type="council")
        page.goto(url)
        wait_for_bundle(page)
        page.wait_for_selector(PROSEMIRROR, timeout=15000)
        page.wait_for_function(f"() => {ALPINE_EDITOR}.collabEnabled === false", timeout=15000)
        page.evaluate(f"() => {ALPINE_EDITOR}.modusSetzen('bearbeiten')")
        tippen(page, " Vor der Zustimmung.")
        with page.expect_navigation(timeout=20000):
            page.get_by_test_id("ablauf-info").get_by_role("button", name="Zustimmen").click()
        anfrage.refresh_from_db()
        assert anfrage.approved is True
        assert "Vor der Zustimmung." in decrypted_content(dokument)
        assert dialoge == [], "nach dem Sichern darf der Browser nicht nach dem Verlassen fragen"
